# -*- coding: utf-8 -*-
"""chart_html -- 数据文件 -> **自包含交互 ECharts HTML**。
蒸馏自 skills.sh 的 smart-charts（只蒸思想，不搬代码）。
支持：.csv / .tsv / .json / .xlsx（openpyxl）。产物零外链（ECharts 已内联）。
"""
from __future__ import annotations
import csv, json, re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ECHARTS = ROOT / "vendor" / "echarts" / "echarts.min.js"


def _read_csv(p: Path, delim: str):
    with p.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=delim))
    if not rows:
        return [], []
    head = [h.strip() or ("c%d" % i) for i, h in enumerate(rows[0])]
    body = [r for r in rows[1:] if any((c or "").strip() for c in r)]
    return head, body


def _read_json(p: Path):
    d = json.loads(p.read_text(encoding="utf-8"))
    if isinstance(d, dict):
        for k in ("data", "rows", "items", "records"):
            if isinstance(d.get(k), list):
                d = d[k]
                break
    if not isinstance(d, list) or not d:
        return [], []
    if isinstance(d[0], dict):
        head = list(d[0].keys())
        body = [[str(r.get(h, "")) for h in head] for r in d]
        return head, body
    head = ["value"]
    return head, [[str(x)] for x in d]


def _read_xlsx(p: Path):
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return [], []
    return [h.strip() for h in rows[0]], rows[1:]


def _num(x):
    s = str(x).strip().replace(",", "").replace("%", "")
    if s in ("", "-", "NA", "N/A", "null", "None"):
        return None
    try:
        return float(s)
    except Exception:
        return None


def load_table(path: str):
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError("数据文件不在：%s" % path)
    ext = p.suffix.lower()
    if ext == ".csv":
        head, body = _read_csv(p, ",")
    elif ext == ".tsv":
        head, body = _read_csv(p, chr(9))
    elif ext == ".json":
        head, body = _read_json(p)
    elif ext in (".xlsx", ".xlsm"):
        head, body = _read_xlsx(p)
    else:
        raise ValueError("不支持的格式：%s（支持 .csv/.tsv/.json/.xlsx）" % ext)
    return head, body


def infer(head, body):
    """推图型与轴：首列非数值且其余数值 ⇒ bar/line；单数值列 ⇒ pie。"""
    if not head or not body:
        return {"kind": "bar", "x": None, "series": []}
    cols = list(range(len(head)))
    numcols = [c for c in cols if sum(1 for r in body[:50] if c < len(r) and _num(r[c]) is not None) >= max(1, int(0.6 * min(50, len(body))))]
    label = [c for c in cols if c not in numcols]
    x = label[0] if label else cols[0]
    series = [c for c in numcols if c != x]
    if not series:
        kind = "pie"
    else:
        vals = [_num(r[series[0]]) for r in body[:20] if series[0] < len(r)]
        vals = [v for v in vals if v is not None]
        looks_date = any(re.search(r"\d{4}[-/]\d{1,2}", str(r[x])) for r in body[:5] if x < len(r))
        kind = "line" if (looks_date or len(body) > 40) else "bar"
    return {"kind": kind, "x": x, "series": series}


def build(head, body, kind=None, title=""):
    info = infer(head, body)
    k = kind or info["kind"]
    x = info["x"]
    xs = [str(r[x]) if x is not None and x < len(r) else str(i) for i, r in enumerate(body)]
    if k == "pie":
        s = info["series"][0] if info["series"] else (info["x"] or 0)
        data = [{"name": xs[i], "value": _num(r[s]) if s < len(r) else None} for i, r in enumerate(body)]
        data = [d for d in data if d["value"] is not None]
        opt = {"title": {"text": title or (head[s] if s < len(head) else "pie")},
               "tooltip": {"trigger": "item"}, "legend": {"top": 28}, "series": [{"type": "pie", "radius": ["35%", "65%"], "data": data}]}
    else:
        series = []
        for c in info["series"]:
            series.append({"name": head[c] if c < len(head) else ("s%d" % c), "type": k, "smooth": k == "line",
                           "data": [_num(r[c]) if c < len(r) else None for r in body]})
        opt = {"title": {"text": title or "chart"},
               "tooltip": {"trigger": "axis"}, "legend": {"top": 28},
               "grid": {"left": 48, "right": 24, "top": 64, "bottom": 48},
               "xAxis": {"type": "category", "data": xs, "axisLabel": {"rotate": 30 if len(xs) > 12 else 0}},
               "yAxis": {"type": "value"}, "series": series}
    return opt, {"kind": k, "rows": len(body), "x": head[x] if (x is not None and x < len(head)) else None,
                 "series": [head[c] for c in info["series"]]}


def run(payload: dict) -> dict:
    op = (payload or {}).get("op") or "make"
    if op != "make":
        return {"ok": False, "status": "user_error", "message": "未知 op：%s（本位只有 make）" % op}
    src = payload.get("data") or payload.get("path")
    if not src:
        return {"ok": False, "status": "user_error", "message": "缺少 data/路径（要转的数据文件）"}
    out = payload.get("out") or (str(ROOT / "state" / "charts" / (Path(src).stem + ".html")))
    out_p = Path(out)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    head, body = load_table(src)
    if not body:
        return {"ok": False, "status": "user_error", "message": "数据为空（读到 0 行）：%s" % src}
    opt, meta = build(head, body, kind=payload.get("kind"), title=payload.get("title") or Path(src).stem)
    lib = ECHARTS.read_text(encoding="utf-8", errors="replace") if ECHARTS.is_file() else ""
    if not lib:
        return {"ok": False, "status": "needs_dependency", "message": "vendor/echarts/echarts.min.js 不在（自包含要它）"}
    html = ("<!doctype html><meta charset=\"utf-8\"><title>" + (payload.get("title") or Path(src).stem)
            + "</title><style>html,body{margin:0;height:100%}#c{width:100%;height:100vh}</style>"
            + "<div id=\"c\"></div><script>" + lib + "</script><script>var o="
            + json.dumps(opt, ensure_ascii=False) + ";var ch=echarts.init(document.getElementById('c'));ch.setOption(o);"
            + "window.addEventListener('resize',function(){ch.resize()});</script>")
    out_p.write_text(html, encoding="utf-8")
    n_pts = sum(len(s.get("data") or []) for s in opt.get("series") or [])
    return {"ok": True, "file": str(out_p), "bytes": out_p.stat().st_size, "kind": meta["kind"],
            "rows": meta["rows"], "series": meta["series"], "points": n_pts,
            "self_contained": all(x not in html for x in ("<script src=", "<link ", "url(http", "src=\"http"))}
